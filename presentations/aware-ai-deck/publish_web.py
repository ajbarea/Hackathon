"""Publish the team deck on the docs site: docs/deck/index.html plus one WebP per slide.

    python publish_web.py AWARE-AI-Hackathon-Deck.pptx <rendered-png-dir>

The PNGs come from rendering the .pptx (graphe's `slides.py render`, PowerPoint or LibreOffice).
Each slide's alt text is its own text; its speaker notes sit in a "What we said" box under it.
Needs python-pptx and pillow.
"""
import html
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "deck"
# Shown under a slide whose figures need context the slide itself does not give.
NOTES = {6: "The new-driver bars (33 % and 32 %) come from the TCN trained without gain augmentation; "
            "with it, the TCN scored 17 % on a new driver. Full tables: examples/TCN.md."}


def slide_text(slide, number: int) -> tuple[str, str]:
    lines = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            lines += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
        if shape.has_table:
            lines += [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
    lines = [x for x in lines if x != str(number)]   # drop the page number
    titles = [sh.text_frame.text.strip() for sh in slide.shapes
              if sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.TITLE]
    title = titles[0] if titles else (lines[0] if lines else "Slide")
    return title, " ".join(x if x[-1] in ".:;?!" else x + "." for x in lines)


def main(pptx: str, pngs: str) -> None:
    deck = Presentation(pptx)
    (OUT / "slides").mkdir(parents=True, exist_ok=True)
    figures = []
    for i, slide in enumerate(deck.slides, 1):
        name = f"slide-{i:02d}.webp"
        im = Image.open(Path(pngs) / f"slide-{i:02d}.png").convert("RGB").resize((1600, 900), Image.LANCZOS)
        im.save(OUT / "slides" / name, "WEBP", quality=88, method=6)
        title, alt = slide_text(slide, i)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        said = (f'<details><summary>What we said</summary><p>{html.escape(notes)}</p></details>'
                if notes else "")
        note = f'<p class="note">{html.escape(NOTES[i])}</p>' if i in NOTES else ""
        figures.append(
            f'<figure id="s{i}"><img src="slides/{name}" width="1600" height="900" loading="{"eager" if i == 1 else "lazy"}"'
            f' alt="{html.escape(alt, quote=True)}"><figcaption><a href="#s{i}">{i}</a> · {html.escape(title)}'
            f'</figcaption>{note}{said}</figure>')
    page = (Path(__file__).with_name("web_deck.html").read_text()
            .replace("{{FIGURES}}", "\n".join(figures)).replace("{{PPTX}}", Path(pptx).name)
            .replace("{{COUNT}}", str(len(figures))))
    (OUT / "index.html").write_text(page)
    (OUT / Path(pptx).name).write_bytes(Path(pptx).read_bytes())
    print(f"{len(figures)} slides -> {OUT}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
